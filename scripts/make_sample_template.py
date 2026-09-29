"""
Generate a sample Bonafide Certificate DOCX template with placeholders.

Run: python scripts/make_sample_template.py
Output: sample_templates/bonafide.docx
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def build_bonafide_template(output_path):
    doc = Document()

    # Header
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('{{ institute_name }}')
    run.bold = True
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = subtitle.add_run('BONAFIDE CERTIFICATE')
    sub.bold = True
    sub.font.size = Pt(14)

    doc.add_paragraph()

    # Body
    body = doc.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.add_run(
        'This is to certify that '
    )
    r = body.add_run('{{ student_name }}')
    r.bold = True
    body.add_run(
        ' (Roll No: {{ roll_no }}), a student of '
        '{{ department }} Department, {{ batch_year }} batch, '
        'is a bona fide student of this institution. '
        'This certificate is issued on {{ today_long }} '
        'for the purpose of {{ purpose }}.'
    )

    doc.add_paragraph()
    doc.add_paragraph()

    # Signatures
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False

    left = table.cell(0, 0).paragraphs[0]
    left.add_run('\n\n_________________\nRegistrar')

    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.add_run('\n\n_________________\nPrincipal')

    # Footer note
    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(
        'This document is digitally signed. Verify authenticity at: {{ verification_url }}'
    )
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f'✓ Template saved: {output_path}')


if __name__ == '__main__':
    build_bonafide_template('sample_templates/bonafide.docx')